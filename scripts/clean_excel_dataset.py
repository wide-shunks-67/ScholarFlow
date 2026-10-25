import openpyxl
import json
import re

excel_path = r'C:\Users\Ansh\.gemini\antigravity\brain\f0a0d8f7-c227-4d79-b05b-6e3208fa5c65\.user_uploaded\media_1792909387743_037ec3fb.xlsx'
wb = openpyxl.load_workbook(excel_path, data_only=True)

INST_MAP = {
    'NIT Arunachal Pradesh': 'NIT Arunachal Pradesh',
    'NIT Tadepalligudem': 'NIT Andhra Pradesh',
    'NIT Suratkal': 'NIT Surathkal',
    'Andhra Pradesh': 'Andhra University & Colleges',
    'IIIT BHAGALPUR': 'IIIT Bhagalpur',
    'BHU': 'Banaras Hindu University (BHU)',
    'MNIT Jaipur': 'MNIT Jaipur',
    'NIT Raipur ': 'NIT Raipur',
    'NIT Goa': 'NIT Goa',
    'NIT Meghalaya': 'NIT Meghalaya',
    'COEP, Pune': 'COEP Tech University, Pune',
    'NIT Kurkshetra': 'NIT Kurukshetra',
    'NIT Srinagar': 'NIT Srinagar',
    'IISER Bhopal': 'IISER Bhopal',
    'NIT Patna': 'NIT Patna',
    'NIT Calicut': 'NIT Calicut',
    'IIT Jodhpur': 'IIT Jodhpur',
    'IIT Madras': 'IIT Madras',
    'IIT Bhubaneshwar': 'IIT Bhubaneswar',
    'Aligarh Muslim University': 'Aligarh Muslim University',
    ' NIT Durgapur': 'NIT Durgapur',
    'NIT Warangal': 'NIT Warangal',
    'IIT Indore': 'IIT Indore',
    'NIT Jalandar, Punjab': 'NIT Jalandhar',
    'NIT Bhopal': 'MANIT Bhopal',
    'IIT Palakkad': 'IIT Palakkad',
    'IIT Dhanbad': 'IIT (ISM) Dhanbad',
    'NIT Trichy': 'NIT Trichy',
    'IIT Ropar': 'IIT Ropar',
    'IIT Bhilai': 'IIT Bhilai',
    'IIT Roorkee': 'IIT Roorkee',
    'NIT Allahabad': 'MNNIT Allahabad',
    'NIT Sikkim': 'NIT Sikkim',
    'NIT Puducheri, United Teriter, ': 'NIT Puducherry',
    'Institute of Engineering & Tech': 'IET Lucknow',
    'IIT Hyderabad': 'IIT Hyderabad',
    'IIT KGP': 'IIT Kharagpur',
    'IIT Goaa': 'IIT Goa',
    'IIT Dharwad': 'IIT Dharwad',
    'IIT Jammu': 'IIT Jammu',
    'IIT Delhi': 'IIT Delhi',
    'IIT Gandhinagar': 'IIT Gandhinagar',
    'IIT Bombay': 'IIT Bombay',
    'IIIT Delhi': 'IIIT Delhi',
    'IIT Kanpur': 'IIT Kanpur',
    'IIT Patna': 'IIT Patna',
    'University of Allahabad': 'University of Allahabad',
    'IIIT Lucknow': 'IIIT Lucknow',
    'IISC Bangalore': 'IISc Bangalore',
    'NIT Uttarakhand': 'NIT Uttarakhand',
    'NIT Manipur': 'NIT Manipur',
    'NIT Delhi': 'NIT Delhi',
    'NIT Hamirpur': 'NIT Hamirpur',
    'RKMVERI': 'RKMVERI Belur',
    'NIT Silchar': 'NIT Silchar',
    'Jadavpur University': 'Jadavpur University',
    'IIIT Bangalore': 'IIIT Bangalore',
    'NIT Surat': 'SVNIT Surat',
    'NIT NAGALAND': 'NIT Nagaland',
    'IIT BHU': 'IIT (BHU) Varanasi',
    'IIITM Gwalior': 'IIITM Gwalior',
    'IGNTU Amarkantak  Cebtral Unive': 'IGNTU Amarkantak',
    'Central University of South Bih': 'Central University of South Bihar',
    'NIT Mizoram': 'NIT Mizoram',
    'IIT Mandi': 'IIT Mandi',
    'IIT Tirupati': 'IIT Tirupati',
    'IIIT Allahabad': 'IIIT Allahabad',
    'IIIT Kota': 'IIIT Kota',
    'IIT Goa': 'IIT Goa',
    'IIT GUWAHATI': 'IIT Guwahati',
    'NIT AGARTALA': 'NIT Agartala',
    'NIT RKL': 'NIT Rourkela',
    'NIT JAIPUR': 'MNIT Jaipur',
    'Sheet56': 'IIEST Shibpur'
}

DEPT_MAP = {
    'cse': 'Computer Science & Engineering',
    'cs': 'Computer Science & Engineering',
    'computer science': 'Computer Science & Engineering',
    'computer science and engineering': 'Computer Science & Engineering',
    'computer science & engineering': 'Computer Science & Engineering',
    'computer engineering': 'Computer Science & Engineering',
    'ece': 'Electronics & Communication Engineering',
    'electronics': 'Electronics & Communication Engineering',
    'electronics and communication': 'Electronics & Communication Engineering',
    'electronics and communication engineering': 'Electronics & Communication Engineering',
    'ee': 'Electrical Engineering',
    'eee': 'Electrical & Electronics Engineering',
    'electrical engineering': 'Electrical Engineering',
    'electrical and electronics engineering': 'Electrical & Electronics Engineering',
    'me': 'Mechanical Engineering',
    'mech': 'Mechanical Engineering',
    'mechanical': 'Mechanical Engineering',
    'mechanical engineering': 'Mechanical Engineering',
    'che': 'Chemical Engineering',
    'chem': 'Chemical Engineering',
    'chemical': 'Chemical Engineering',
    'chemical engineering': 'Chemical Engineering',
    'it': 'Information Technology',
    'information technology': 'Information Technology',
    'math': 'Mathematics & Computing',
    'maths': 'Mathematics & Computing',
    'mathematics': 'Mathematics & Computing',
    'maths & computing': 'Mathematics & Computing',
    'physics': 'Physics & Applied Sciences',
    'applied physics': 'Physics & Applied Sciences',
    'chemistry': 'Chemistry & Chemical Sciences'
}

def clean_dept(raw):
    if not raw or raw.lower() in ['none', 'nan', '-']:
        return 'Engineering & Sciences'
    s = raw.strip()
    s_lower = s.lower().replace('.', '')
    return DEPT_MAP.get(s_lower, s)

all_profs = []
seen_emails = set()

for sheetname in wb.sheetnames:
    norm_inst = INST_MAP.get(sheetname)
    if not norm_inst:
        continue
    ws = wb[sheetname]
    
    header_row = 1
    name_col = None
    email_col = None
    dept_col = None
    area_col = None
    
    for r in range(1, min(6, ws.max_row + 1)):
        row_vals = [str(ws.cell(r, c).value or '').strip() for c in range(1, ws.max_column + 1)]
        row_lower = [v.lower() for v in row_vals]
        for idx, val in enumerate(row_lower):
            if any(k in val for k in ['name of the professor', 'name of the faculty', 'name', 'faculty name']) and not name_col:
                name_col = idx + 1
            if ('email' in val or 'mail' in val) and not email_col:
                email_col = idx + 1
            if any(k in val for k in ['dept', 'department', 'centre']) and not dept_col:
                dept_col = idx + 1
            if any(k in val for k in ['area', 'research', 'interest', 'specialization']) and not area_col:
                area_col = idx + 1
        if name_col and email_col:
            header_row = r
            break
            
    if not (name_col and email_col):
        name_col = 2
        email_col = 3
        dept_col = 4
        area_col = 5
        header_row = 1

    start_r = header_row + 1
    for r in range(start_r, ws.max_row + 1):
        name = ws.cell(r, name_col).value if name_col else None
        email = ws.cell(r, email_col).value if email_col else None
        dept = ws.cell(r, dept_col).value if dept_col else ''
        area = ws.cell(r, area_col).value if area_col else ''
        
        name = str(name).strip() if name is not None else ''
        email = str(email).strip() if email is not None else ''
        dept = str(dept).strip() if dept is not None else ''
        area = str(area).strip() if area is not None else ''
        
        if len(name) < 3 or name.lower() in ['none', 'nan', 'name', 'faculty', 'professor']:
            continue
            
        if '@' not in email:
            # check other columns in case email was shifted
            found = False
            for c in range(1, ws.max_column + 1):
                val = str(ws.cell(r, c).value or '').strip()
                if '@' in val and '.' in val:
                    email = val
                    found = True
                    break
            if not found:
                continue

        # Extract primary email cleanly
        emails = [e.strip() for e in re.split(r'[,;/\s]+', email) if '@' in e and '.' in e]
        if not emails:
            continue
        primary_email = emails[0].lower().replace('[at]', '@').replace('[dot]', '.').replace('(at)', '@').replace('(dot)', '.')
        
        # Deduplicate
        if primary_email in seen_emails:
            continue
        seen_emails.add(primary_email)
        
        # Clean clean name
        clean_name = re.sub(r'\s+', ' ', name).strip()
        clean_dept_name = clean_dept(dept)
        clean_area = re.sub(r'\s+', ' ', area).strip()
        if not clean_area or clean_area.lower() in ['none', 'nan', '-', 'nil']:
            clean_area = f"Research & Engineering in {clean_dept_name}"
            
        all_profs.append({
            'id': f'prof-{len(all_profs)+1}',
            'name': clean_name,
            'institution': norm_inst,
            'department': clean_dept_name,
            'researchInterests': clean_area,
            'email': primary_email
        })

print(f"Total verified unique professors from user Excel: {len(all_profs)}")
print(f"Total unique institutions: {len(set(p['institution'] for p in all_profs))}")
print(f"Total unique departments: {len(set(p['department'] for p in all_profs))}")

# Write cleanly to frontend/professors.json and frontend/professors_data.js
output_json = r'C:\Users\Ansh\OneDrive\Desktop\Hacktoberfest\scholarflow\frontend\professors.json'
output_js = r'C:\Users\Ansh\OneDrive\Desktop\Hacktoberfest\scholarflow\frontend\professors_data.js'

with open(output_json, 'w', encoding='utf-8') as f:
    json.dump(all_profs, f, indent=2)

with open(output_js, 'w', encoding='utf-8') as f:
    f.write('window.SCHOLARFLOW_PROFESSORS = ' + json.dumps(all_profs) + ';\n')

print(f"Saved cleanly to {output_json} and {output_js}!")

